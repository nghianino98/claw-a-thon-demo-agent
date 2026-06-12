from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import sqlite3
import subprocess
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.db import Database
from app.services.audit import AuditService
from app.services.registry import SkillRegistry, WorkflowRegistry
from app.settings import Settings
from app.utils import compact_text, log_event, safe_json, sha256_file, utc_now


TEXT_EXTENSIONS = {
    ".md",
    ".markdown",
    ".txt",
    ".csv",
    ".json",
    ".yaml",
    ".yml",
    ".html",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".go",
    ".sql",
    ".gs",
    ".sh",
    ".pdf",
    ".docx",
    ".xlsx",
    ".xlsm",
}

EXCLUDED_DIRS = {".git", ".next", ".obsidian", ".vscode", ".sixth", "venv", "node_modules", "__pycache__"}
EXCLUDED_NAMES = {".DS_Store", ".greennode.json"}
PRODUCTS = ["MMF", "FD", "FI", "CCQ", "Insurance", "Stock", "Crypto", "FS Hub", "FS Profile"]
PRODUCT_ALIASES = {
    "MMF": ["mmf", "money market"],
    "FD": ["fd", "fixed deposit"],
    "FI": ["fi", "fixed income", "trái phiếu", "trai phieu"],
    "CCQ": ["ccq", "chứng chỉ quỹ", "chung chi quy", "fund"],
    "Insurance": ["insurance", "bảo hiểm", "bao hiem"],
    "Stock": ["stock", "chứng khoán", "chung khoan"],
    "Crypto": ["crypto"],
    "FS Hub": ["fs hub"],
    "FS Profile": ["fs profile", "fs profle"],
}


@dataclass
class KnowledgeHit:
    path: str
    title: str
    snippet: str
    score: float
    lines: str
    product: str
    area: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "title": self.title,
            "snippet": self.snippet,
            "score": self.score,
            "lines": self.lines,
            "product": self.product,
            "area": self.area,
        }


class KBService:
    def __init__(
        self,
        db: Database,
        settings: Settings,
        audit: AuditService | None = None,
        skill_registry: SkillRegistry | None = None,
        workflow_registry: WorkflowRegistry | None = None,
        telegram_client: Any | None = None,
    ):
        self.db = db
        self.settings = settings
        self.audit = audit
        self.skill_registry = skill_registry
        self.workflow_registry = workflow_registry
        self.telegram_client = telegram_client
        self.kb_root = settings.kb_dir
        self.versions_dir = self.kb_root / "versions"
        self.current_link = self.kb_root / "current"

    def ensure_dirs(self) -> None:
        self.versions_dir.mkdir(parents=True, exist_ok=True)
        self.settings.artifacts_dir.mkdir(parents=True, exist_ok=True)
        (self.settings.state_dir / "uploads" / "tmp").mkdir(parents=True, exist_ok=True)
        (self.settings.state_dir / "backups").mkdir(parents=True, exist_ok=True)

    def active_version(self) -> int | None:
        with self.db.connect() as conn:
            row = conn.execute("SELECT id FROM kb_versions WHERE status='active' ORDER BY id DESC LIMIT 1").fetchone()
        return int(row["id"]) if row else None

    def status(self) -> dict[str, Any]:
        version = self.active_version()
        with self.db.connect() as conn:
            row = conn.execute(
                "SELECT * FROM kb_versions WHERE id=?",
                (version,),
            ).fetchone() if version else None
        return {
            "available": bool(version and self.current_link.exists()),
            "kb_version": version,
            "version": dict(row) if row else None,
            "current_path": str(self.current_link),
        }

    def build_from_directory(self, source: Path, uploaded_by: str = "cli", activate: bool = True) -> int:
        self.ensure_dirs()
        if not source.exists() or not source.is_dir():
            raise FileNotFoundError(str(source))
        version_id = self._create_version(uploaded_by, source.name)
        target = self.versions_dir / str(version_id)
        try:
            self._copy_tree(source, target)
            stats = self._index_version(version_id, target)
            skills = workflows = 0
            if activate:
                self.activate_version(version_id)
                skills, workflows = self.reload_registries()
            with self.db.connect() as conn:
                conn.execute(
                    """
                    UPDATE kb_versions
                    SET status=?, file_count=?, chunk_count=?, total_bytes=?, skill_count=?, workflow_count=?
                    WHERE id=?
                    """,
                    ("active" if activate else "ready", stats["files"], stats["chunks"], stats["bytes"], skills, workflows, version_id),
                )
                conn.commit()
            if self.audit:
                self.audit.record(uploaded_by, "kb_build", str(version_id), stats)
            return version_id
        except Exception as exc:
            with self.db.connect() as conn:
                conn.execute("UPDATE kb_versions SET status='failed', error=? WHERE id=?", (str(exc), version_id))
                conn.commit()
            raise

    def build_from_zip(self, archive_path: Path, uploaded_by: str = "admin-api", activate: bool = True) -> int:
        self.ensure_dirs()
        if not archive_path.exists() or not zipfile.is_zipfile(archive_path):
            raise ValueError("Upload must be a valid zip file")
        max_compressed = self.settings.kb_upload_max_mb * 1024 * 1024
        if archive_path.stat().st_size > max_compressed:
            raise ValueError("Zip exceeds KB_UPLOAD_MAX_MB")
        version_id = self._create_version(uploaded_by, archive_path.name)
        target = self.versions_dir / str(version_id)
        try:
            self._extract_zip(archive_path, target)
            stats = self._index_version(version_id, target)
            skills = workflows = 0
            if activate:
                self.activate_version(version_id)
                skills, workflows = self.reload_registries()
            with self.db.connect() as conn:
                conn.execute(
                    """
                    UPDATE kb_versions
                    SET status=?, file_count=?, chunk_count=?, total_bytes=?, skill_count=?, workflow_count=?
                    WHERE id=?
                    """,
                    ("active" if activate else "ready", stats["files"], stats["chunks"], stats["bytes"], skills, workflows, version_id),
                )
                conn.commit()
            if self.audit:
                self.audit.record(uploaded_by, "kb_upload", str(version_id), stats)
            return version_id
        except Exception as exc:
            with self.db.connect() as conn:
                conn.execute("UPDATE kb_versions SET status='failed', error=? WHERE id=?", (str(exc), version_id))
                conn.commit()
            raise

    def activate_version(self, version_id: int) -> None:
        version_path = self.versions_dir / str(version_id)
        if not version_path.exists():
            raise FileNotFoundError(str(version_path))
        tmp = self.kb_root / f".current-{version_id}.tmp"
        if tmp.exists() or tmp.is_symlink():
            tmp.unlink()
        relative_target = os.path.relpath(version_path, tmp.parent)
        os.symlink(relative_target, tmp)
        os.replace(tmp, self.current_link)
        now = utc_now()
        with self.db.connect() as conn:
            conn.execute("UPDATE kb_versions SET status='archived' WHERE status='active' AND id<>?", (version_id,))
            conn.execute("UPDATE kb_versions SET status='active', activated_at=? WHERE id=?", (now, version_id))
            conn.commit()
        self.reload_registries()
        self._refresh_kb_map_cache()
        self._cleanup_old_versions()
        if self.audit:
            self.audit.record("system", "kb_activate", str(version_id), {})

    def reload_registries(self) -> tuple[int, int]:
        skills = self.skill_registry.reload_from_kb() if self.skill_registry else 0
        workflows = self.workflow_registry.reload_from_kb() if self.workflow_registry else 0
        return skills, workflows

    def _refresh_kb_map_cache(self) -> None:
        cache = self._build_kb_map_cache()
        with self.db.connect() as conn:
            conn.execute(
                """
                INSERT INTO settings(key, value, updated_at, updated_by)
                VALUES ('kb_map_cache', ?, ?, 'system')
                ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=excluded.updated_at, updated_by=excluded.updated_by
                """,
                (cache, utc_now()),
            )
            conn.commit()

    def _build_kb_map_cache(self) -> str:
        root = self.current_link
        if not root.exists():
            return "(không tìm thấy KB hiện tại)"
        entries: list[str] = []
        try:
            resolved = root.resolve()
            for dirname in ["05. Knowledge", "01. Objective", "04. Skill", "02. Context", "03. Fact"]:
                base = resolved / dirname
                if not base.exists():
                    continue
                for item in sorted(base.rglob("*")):
                    if len(entries) >= 200:
                        break
                    if item.name.startswith("."):
                        continue
                    if "Zalopay Design System" in item.parts or "Tokens" in item.parts:
                        continue
                    if item.suffix.lower() in {".xlsx", ".pdf", ".docx", ".pptx", ".zip", ".json"}:
                        continue
                    rel = item.relative_to(resolved)
                    entries.append(rel.as_posix() + ("/" if item.is_dir() else ""))
            if len(entries) >= 200:
                entries.append("[...] (còn nhiều file context/fact khác)")
        except Exception:
            return "(lỗi khi dựng bản đồ KB)"
        return "\n".join(entries) if entries else "(KB hiện tại chưa có file phù hợp để lập bản đồ)"

    def versions(self) -> list[dict[str, Any]]:
        with self.db.connect() as conn:
            rows = conn.execute("SELECT * FROM kb_versions ORDER BY id DESC").fetchall()
        return [dict(row) for row in rows]

    def manifest(self) -> dict[str, Any]:
        version = self.active_version()
        if not version:
            return {"kb_version": None, "files": {}}
        with self.db.connect() as conn:
            rows = conn.execute("SELECT path, sha256, size FROM kb_files WHERE kb_version=? ORDER BY path", (version,)).fetchall()
        return {
            "kb_version": version,
            "files": {row["path"]: {"sha256": row["sha256"], "size": row["size"]} for row in rows},
        }

    def search(self, query: str, product: str | None = None, area: str | None = None, top_k: int = 5) -> list[KnowledgeHit]:
        version = self.active_version()
        if not version:
            return []
        top_k = max(1, min(top_k, 10))
        match = self._match_query(query)
        if not match:
            return []
        rows = self._search_sql(version, match, product, area, top_k * 5)
        if not rows and (product or area):
            rows = self._search_sql(version, match, None, None, top_k * 5)
        hits = [self._row_to_hit(row, query) for row in rows]
        hits.sort(key=lambda hit: hit.score)
        return hits[:top_k]

    async def search_async(self, query: str, product: str | None = None, area: str | None = None, top_k: int = 5) -> list[KnowledgeHit]:
        version = self.active_version()
        if not version:
            return []

        phrase_q, and_q, or_q = self._compile_fts_queries(query)
        if not phrase_q:
            return []

        top_k = max(1, min(top_k, 10))
        limit = 200


        import asyncio
        rows_phrase, rows_and, rows_or = await asyncio.gather(
            asyncio.to_thread(self._search_sql, version, f'"{phrase_q}"', product, area, limit),
            asyncio.to_thread(self._search_sql, version, and_q, product, area, limit),
            asyncio.to_thread(self._search_sql, version, or_q, product, area, limit)
        )
        rows_hint = await asyncio.to_thread(self._path_hint_sql, version, query, product, area, limit=40)

        row_map = {}
        rrf_scores = {}
        k_rrf = 60

        def process_list(rows):
            for rank, row in enumerate(rows, start=1):
                rid = row["rowid"]
                row_map[rid] = row
                rrf_scores[rid] = rrf_scores.get(rid, 0.0) + (1.0 / (k_rrf + rank))

        process_list(rows_phrase)
        process_list(rows_and)
        process_list(rows_or)
        for row in rows_hint:
            rid = row["rowid"]
            row_map[rid] = row
            rrf_scores[rid] = rrf_scores.get(rid, 0.0) + 0.08

        hits = []
        from datetime import datetime, UTC
        for rid, rrf in rrf_scores.items():
            row = row_map[rid]
            days_old = 0
            mtime_str = None
            try:
                mtime_str = row["mtime"]
            except (KeyError, IndexError):
                pass
            if mtime_str:
                try:
                    file_dt = datetime.fromisoformat(mtime_str.replace("Z", "+00:00"))
                    days_old = (datetime.now(UTC) - file_dt).days
                except Exception:
                    pass

            penalty = 0.0
            if days_old > 30:
                penalty = min(3.0, (days_old - 30) / 30.0)

            rrf_penalty = (penalty / 3.0) * 0.02
            score = -rrf + rrf_penalty

            haystack = f"{row['path']} {row['title']}".lower()
            query_l = query.lower()
            for prod, aliases in PRODUCT_ALIASES.items():
                if any(alias in query_l for alias in aliases) and any(alias in haystack for alias in aliases):
                    score -= 0.15
            if row["area"] == "fact" and re.search(r"\b(lỗi|bug|issue|ticket|sự cố|error|fail)\b", query_l):
                score -= 0.02
            score += self._path_score_adjustment(query, row["path"], row["title"], row["area"])
            if row["area"] == "knowledge" and not self._has_non_knowledge_intent(query):
                score -= 0.25

            snippet = compact_text(str(row["content"]))[:400]
            hit = KnowledgeHit(
                path=row["path"],
                title=row["title"],
                snippet=snippet,
                score=score,
                lines=f"L{row['start_line']}-L{row['end_line']}",
                product=row["product"],
                area=row["area"],
            )
            hits.append(hit)

        hits.sort(key=lambda x: x.score)

        grouped_hits = []
        path_counts = {}
        for hit in hits:
            path_counts[hit.path] = path_counts.get(hit.path, 0) + 1
            if path_counts[hit.path] <= 2:
                grouped_hits.append(hit)

        return grouped_hits[:top_k]

    @staticmethod
    def _compile_fts_queries(query: str) -> tuple[str, str, str]:
        stopwords = {
            "và", "của", "là", "thì", "mà", "nhưng", "cũng", "các", "những", "một", "cho", "để", "với", "tại", "trong",
            "ra", "vào", "lên", "xuống", "về", "đến", "theo", "qua", "bởi", "vì", "nên", "nếu", "tuy", "được", "bị", "này", "kia", "đó",
            "the", "and", "of", "is", "in", "on", "at", "to", "for", "with", "by", "an", "a", "that", "this", "these", "those", "it", "its",
            "đang", "có", "không", "ở", "từ", "trên", "dưới", "ngoài", "loại", "cho", "làm", "thế", "nào", "sao", "gì", "đâu", "ai", "mình",
            "bạn", "này", "đó", "hoạt", "động", "cơ", "chế", "luồng", "quy", "trình", "bước", "lịch", "trình", "thời", "gian", "ngày",
            "tháng", "năm", "giúp", "nhé", "nha", "đi", "hỏi", "vấn", "đề", "sự", "cố", "chi", "tiết", "cho", "biết", "hiện", "tại",
            "bán", "mua", "gói", "sản", "phẩm", "tìm", "hiểu", "thông", "tin", "xử", "lý"
        }
        mappings = {
            "nạp tiền": ["deposit"],
            "nap tien": ["deposit"],
            "rút tiền": ["withdraw", "redemption", "redeem"],
            "rut tien": ["withdraw", "redemption", "redeem"],
            "thanh toán": ["payment"],
            "thanh toan": ["payment"],
            "đối soát": ["reconciliation", "reconcile"],
            "doi soat": ["reconciliation", "reconcile"],
            "chuyển tiền": ["transfer"],
            "chuyen tien": ["transfer"],
            "tài khoản": ["account"],
            "tai khoan": ["account"],
            "số dư": ["balance"],
            "so du": ["balance"],
            "kỹ thuật": ["controller", "handler", "service", "internal"],
            "ky thuat": ["controller", "handler", "service", "internal"],
            "mã nguồn": ["controller", "handler", "service", "internal"],
            "ma nguon": ["controller", "handler", "service", "internal"],
            "code": ["controller", "handler", "service", "internal"],
            "lỗi": ["error", "fail", "failed", "pending"],
            "loi": ["error", "fail", "failed", "pending"],
            "thất bại": ["fail", "failed"],
            "that bai": ["fail", "failed"],
        }
        lower_query = query.lower()
        tokens_with_pos = []
        for m in re.finditer(r"[\wÀ-ỹ]+", lower_query, flags=re.UNICODE):
            token = m.group(0)
            pos = m.start()
            tokens_with_pos.append((pos, token))

        raw_words = [t for _, t in tokens_with_pos if t not in stopwords and len(t) >= 2]
        if not raw_words:
            raw_words = [t for _, t in tokens_with_pos if len(t) >= 2]

        if not raw_words:
            return "", "", ""

        extra = []
        sorted_keys = sorted(mappings.keys(), key=len, reverse=True)
        for key in sorted_keys:
            if key in lower_query:
                extra.extend(mappings[key])

        phrase_query = " ".join(raw_words)

        and_clauses = []
        used_keys = []
        for key in sorted_keys:
            if key in lower_query:
                already_covered = False
                for uk in used_keys:
                    if key in uk or uk in key:
                        already_covered = True
                        break
                if not already_covered:
                    used_keys.append(key)
                    words = key.split()
                    eng_terms = mappings[key]
                    clause = f"({' AND '.join(words)} OR {' OR '.join(eng_terms)})"
                    and_clauses.append(clause)

        for word in raw_words:
            is_part_of_key = False
            for uk in used_keys:
                if word in uk.split():
                    is_part_of_key = True
                    break
            if not is_part_of_key:
                if word in mappings:
                    eng_terms = mappings[word]
                    and_clauses.append(f"({word} OR {' OR '.join(eng_terms)})")
                else:
                    and_clauses.append(word)

        and_query = " AND ".join(and_clauses)
        or_terms = list(set(raw_words + extra))
        or_query = " OR ".join(or_terms)

        return phrase_query, and_query, or_query

    def read(self, path: str, start_line: int | None = None, end_line: int | None = None, max_chars: int | None = None) -> str:
        file_path = self._safe_current_path(path)
        if not file_path.exists():
            return f'{{"error":"not found: {path}"}}'
        if file_path.suffix.lower() not in TEXT_EXTENSIONS:
            return f'{{"error":"file is not text-indexable in v1: {path}"}}'
        lines = self._read_text(file_path).splitlines()
        total = len(lines)
        start = max(1, start_line or 1)
        end = min(total, end_line or total)
        selected = lines[start - 1 : end]
        text = "\n".join(f"{idx}: {line}" for idx, line in enumerate(selected, start=start))
        limit = max_chars or self.settings.kb_read_max_chars
        if len(text) > limit:
            text = text[:limit] + f"\n[đã cắt bớt, đọc tiếp với start_line={start + len(selected)}]"
        return text

    def list_tree(self, path: str = ".", depth: int = 2) -> str:
        root = self._safe_current_path(path)
        if not root.exists() or not root.is_dir():
            return f'{{"error":"directory not found: {path}"}}'
        depth = max(0, min(depth, 2))
        entries: list[str] = []
        base_depth = len(root.relative_to(self.current_link.resolve()).parts) if root != self.current_link.resolve() else 0
        for item in sorted(root.rglob("*")):
            rel = item.relative_to(self.current_link.resolve())
            current_depth = len(rel.parts) - base_depth
            if current_depth > depth + 1:
                continue
            entries.append(rel.as_posix() + ("/" if item.is_dir() else ""))
            if len(entries) >= 200:
                entries.append("[đã cắt bớt sau 200 entries]")
                break
        return "\n".join(entries)

    def grep(self, pattern: str, path_prefix: str | None = None, max_results: int = 30) -> str:
        max_results = max(1, min(max_results, 50))
        root = self._safe_current_path(path_prefix or ".")
        if not root.exists():
            return f'{{"error":"path_prefix not found: {path_prefix}"}}'
        rg = shutil.which("rg")
        if rg:
            return self._grep_rg(rg, pattern, root, max_results)
        return self._grep_python(pattern, root, max_results)

    def _create_version(self, uploaded_by: str, original_filename: str) -> int:
        with self.db.connect() as conn:
            cur = conn.execute(
                """
                INSERT INTO kb_versions(status, uploaded_by, original_filename, created_at, kind)
                VALUES ('processing', ?, ?, ?, 'full')
                """,
                (uploaded_by, original_filename, utc_now()),
            )
            conn.commit()
            return int(cur.lastrowid)

    def _copy_tree(self, source: Path, target: Path) -> None:
        if target.exists():
            shutil.rmtree(target)

        def copy_function(src: str, dst: str) -> str:
            try:
                os.link(src, dst)
            except OSError:
                shutil.copy2(src, dst)
            return dst

        shutil.copytree(source, target, ignore=self._ignore_names, copy_function=copy_function, symlinks=False)

    def _extract_zip(self, archive_path: Path, target: Path) -> None:
        if target.exists():
            shutil.rmtree(target)
        target.mkdir(parents=True, exist_ok=True)
        max_uncompressed = self.settings.kb_upload_max_mb * 1024 * 1024 * 5
        max_entries = 80000
        total = 0
        with zipfile.ZipFile(archive_path) as zf:
            infos = zf.infolist()
            if len(infos) > max_entries:
                raise ValueError("Zip has too many entries")
            for info in infos:
                rel = Path(info.filename)
                if info.is_dir():
                    continue
                if rel.is_absolute() or ".." in rel.parts:
                    raise ValueError("Zip path traversal blocked")
                if self._excluded_zip_path(rel):
                    continue
                total += info.file_size
                if total > max_uncompressed:
                    raise ValueError("Zip uncompressed size exceeds safety limit")
                mode = (info.external_attr >> 16) & 0o170000
                if mode == 0o120000:
                    raise ValueError("Zip symlink entries are not allowed")
                destination = (target / rel).resolve()
                try:
                    destination.relative_to(target.resolve())
                except ValueError as exc:
                    raise ValueError("Zip path traversal blocked") from exc
                destination.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(info, "r") as src, destination.open("wb") as dst:
                    shutil.copyfileobj(src, dst, length=1024 * 1024)

    @staticmethod
    def _excluded_zip_path(rel: Path) -> bool:
        parts = rel.parts
        lowered = [part.lower() for part in parts]
        if any(part in EXCLUDED_DIRS for part in parts):
            return True
        name = parts[-1]
        low_name = name.lower()
        return (
            name in EXCLUDED_NAMES
            or low_name.endswith(".pyc")
            or low_name.endswith(".sqlite3")
            or ".env" in low_name
            or "credentials" in low_name
        )

    def _ignore_names(self, directory: str, names: list[str]) -> set[str]:
        ignored: set[str] = set()
        for name in names:
            path = Path(directory) / name
            lower = name.lower()
            if name in EXCLUDED_NAMES or lower in EXCLUDED_DIRS:
                ignored.add(name)
            elif lower.endswith(".pyc") or ".env" in lower or "credentials" in lower or lower.endswith(".sqlite3"):
                ignored.add(name)
            elif path.is_symlink():
                ignored.add(name)
        return ignored

    def _index_version(self, version_id: int, root: Path) -> dict[str, int]:
        stats = {"files": 0, "chunks": 0, "bytes": 0}
        with self.db.connect() as conn:
            for file_path in self._iter_files(root):
                rel = file_path.relative_to(root).as_posix()
                size = file_path.stat().st_size
                sha = sha256_file(file_path)
                mtime = utc_now()
                conn.execute(
                    """
                    INSERT OR REPLACE INTO kb_files(kb_version, path, sha256, size, mtime)
                    VALUES (?,?,?,?,?)
                    """,
                    (version_id, rel, sha, size, mtime),
                )
                stats["files"] += 1
                stats["bytes"] += size
                if file_path.suffix.lower() not in TEXT_EXTENSIONS:
                    continue
                if size > self.settings.kb_file_max_mb * 1024 * 1024:
                    continue
                exists = conn.execute(
                    "SELECT 1 FROM chunks_meta WHERE path=? AND file_sha=? LIMIT 1",
                    (rel, sha),
                ).fetchone()
                if exists:
                    continue
                text = self._read_text(file_path)
                title = self._title_for(rel, text)
                product = self._classify_product(rel)
                area = self._classify_area(rel)
                for chunk, start_line, end_line in self._chunk_text(text, rel):
                    cur = conn.execute(
                        "INSERT INTO chunks_fts(path, title, product, area, content) VALUES (?,?,?,?,?)",
                        (rel, title, product, area, chunk),
                    )
                    conn.execute(
                        "INSERT INTO chunks_meta(rowid_fts, path, file_sha, start_line, end_line, mtime) VALUES (?,?,?,?,?,?)",
                        (cur.lastrowid, rel, sha, start_line, end_line, mtime),
                    )
                    stats["chunks"] += 1
            conn.commit()
        return stats

    def _iter_files(self, root: Path):
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [name for name in dirnames if name not in EXCLUDED_DIRS and not name.startswith(".git")]
            for filename in filenames:
                if filename in EXCLUDED_NAMES or filename.endswith(".pyc"):
                    continue
                lower = filename.lower()
                if ".env" in lower or "credentials" in lower or lower.endswith(".sqlite3"):
                    continue
                path = Path(dirpath) / filename
                if path.is_symlink() or not path.is_file():
                    continue
                yield path

    def _search_sql(
        self, version: int, match: str, product: str | None, area: str | None, limit: int
    ) -> list[sqlite3.Row]:
        where = ["chunks_fts MATCH ?", "kf.kb_version=?"]
        params: list[Any] = [match, version]
        if product:
            where.append("chunks_fts.product=?")
            params.append(product)
        if area:
            where.append("chunks_fts.area=?")
            params.append(area)
        params.append(limit)
        sql = f"""
            SELECT chunks_fts.rowid, chunks_fts.path, chunks_fts.title, chunks_fts.product, chunks_fts.area,
                   chunks_fts.content, bm25(chunks_fts) AS score, m.start_line, m.end_line, m.mtime
            FROM chunks_fts
            CROSS JOIN chunks_meta m ON chunks_fts.rowid=m.rowid_fts
            CROSS JOIN kb_files kf ON kf.path=m.path AND kf.sha256=m.file_sha
            WHERE {' AND '.join(where)}
            ORDER BY (CASE
                WHEN chunks_fts.area = 'knowledge' THEN bm25(chunks_fts) - 5.0
                WHEN chunks_fts.area = 'objective' THEN bm25(chunks_fts) - 3.0
                ELSE bm25(chunks_fts)
            END)
            LIMIT ?
        """
        with self.db.connect() as conn:
            return conn.execute(sql, tuple(params)).fetchall()

    def _path_hint_sql(
        self,
        version: int,
        query: str,
        product: str | None,
        area: str | None,
        limit: int = 40,
    ) -> list[sqlite3.Row]:
        patterns = self._path_hint_patterns(query)
        if not patterns:
            return []
        where = ["kf.kb_version=?"]
        params: list[Any] = [version]
        if product:
            where.append("chunks_fts.product=?")
            params.append(product)
        if area:
            where.append("chunks_fts.area=?")
            params.append(area)
        pattern_clauses = []
        for pattern in patterns[:12]:
            pattern_clauses.append("m.path LIKE ?")
            params.append(pattern)
        where.append("(" + " OR ".join(pattern_clauses) + ")")
        params.append(limit)
        sql = f"""
            SELECT chunks_fts.rowid, chunks_fts.path, chunks_fts.title, chunks_fts.product, chunks_fts.area,
                   chunks_fts.content, 0.0 AS score, m.start_line, m.end_line, m.mtime
            FROM chunks_fts
            CROSS JOIN chunks_meta m ON chunks_fts.rowid=m.rowid_fts
            CROSS JOIN kb_files kf ON kf.path=m.path AND kf.sha256=m.file_sha
            WHERE {' AND '.join(where)}
            ORDER BY m.path, m.start_line
            LIMIT ?
        """
        with self.db.connect() as conn:
            return conn.execute(sql, tuple(params)).fetchall()

    @classmethod
    def _path_hint_patterns(cls, query: str) -> list[str]:
        q = cls._fold(query)
        query_products = cls._query_products(query)
        patterns: list[str] = []
        if cls._contains_phrase(q, "wealth") and cls._has_any(q, ["san pham", "products", "liet ke", "hien co"]):
            if not cls._has_any(q, ["strategy", "chien luoc", "roadmap", "ap26"]):
                patterns.append("05. Knowledge/_Index.md")
        if cls._has_any(q, ["strategy", "chien luoc", "roadmap", "uu tien", "priority", "ap26", "focus"]):
            patterns.extend(
                [
                    "01. Objective/Product Strategy/%",
                    "02. Context/Confluence/Wealth General/%Planning%",
                    "02. Context/Confluence/Wealth General/%Roadmap%",
                    "02. Context/Confluence/Wealth General/%Plan%",
                ]
            )
        if cls._has_any(q, ["monthly report", "highlight", "ship", "thang 2", "thang 3", "thang 4", "thang 5", "feb", "mar", "apr", "may"]):
            month_patterns = {
                "thang 2": ["01. Objective/Monthly Report/%Feb Report%"],
                "feb": ["01. Objective/Monthly Report/%Feb Report%"],
                "thang 3": ["01. Objective/Monthly Report/%Mar Report%"],
                "mar": ["01. Objective/Monthly Report/%Mar Report%"],
                "thang 4": ["01. Objective/Monthly Report/%Apr Report%"],
                "apr": ["01. Objective/Monthly Report/%Apr Report%"],
                "thang 5": ["01. Objective/Monthly Report/%May Report%"],
                "may": ["01. Objective/Monthly Report/%May Report%"],
            }
            added = False
            for marker, vals in month_patterns.items():
                if cls._contains_phrase(q, marker):
                    patterns.extend(vals)
                    added = True
            if not added:
                patterns.append("01. Objective/Monthly Report/%")
        if cls._has_any(q, ["kpi", "okr", "objective", "kr", "product contribution", "shared kpi", "chi so"]):
            patterns.extend(["01. Objective/Product KPI/%"])
        if cls._contains_phrase(q, "product contribution"):
            patterns.append("01. Objective/Product KPI/01. Product Contribution/%")
        if cls._contains_phrase(q, "shared kpi"):
            patterns.append("01. Objective/Product KPI/02. Shared KPI/%")
        if cls._has_any(q, ["checklist", "audit", "financial safety", "trust-building", "trust building"]):
            patterns.extend(
                [
                    ".agents/skills/Audit/Product Audit/%",
                    "04. Skill/%Audit%",
                    "01. Objective/Product KPI/03. Product Audit/%Checklist%",
                    "01. Objective/Product KPI/03. Product Audit/%checklist%",
                ]
            )
        exact_issue = re.findall(r"\bISSUE-\d+\b", query, flags=re.IGNORECASE)
        for issue in exact_issue:
            patterns.append(f"03. Fact/CS Ticket/%{issue}%")
        ticket_codes = re.findall(r"\b[A-Z]{2,10}-\d+\b", query)
        for code in ticket_codes:
            patterns.append(f"02. Context/Jira/%{code}%")
        if not exact_issue and cls._has_any(q, ["ticket", "issue", "khieu nai", "cs ticket"]):
            patterns.append("03. Fact/CS Ticket/%")
        if cls._has_any(q, ["transid", "transaction", "timeout", "retry"]):
            patterns.append("03. Fact/Issue Investigation/%")
        if cls._has_any(q, ["source code", "ma nguon", "code", "ham", "function", "logic", "file nao", "controller", "handler"]):
            if cls._contains_phrase(q, "redemption"):
                if query_products:
                    for product in query_products:
                        patterns.extend(
                            [
                                f"03. Fact/Source Code/{product}/%redemption%",
                                f"03. Fact/Source Code/{product}/Backend/redemption/%",
                                f"03. Fact/Source Code/{product}/Backend/order-worker/handler/%redemption%",
                                f"03. Fact/Source Code/{product}/Frontend/%Redemption%",
                            ]
                        )
                else:
                    patterns.append("03. Fact/Source Code/%redemption%")
            else:
                if query_products:
                    for product in query_products:
                        patterns.extend(
                            [
                                f"03. Fact/Source Code/{product}/%controller%",
                                f"03. Fact/Source Code/{product}/%handler%",
                            ]
                        )
                patterns.extend(["03. Fact/Source Code/%controller%", "03. Fact/Source Code/%handler%"])
        if cls._has_any(q, ["fs profile", "kyc", "risk assessment"]):
            patterns.extend(["02. Context/Confluence/FS Profile/%", "02. Context/Confluence/FS Profle/%", "02. Context/Confluence/FS Hub/%Profile%"])
        deduped: list[str] = []
        for pattern in patterns:
            if pattern not in deduped:
                deduped.append(pattern)
        return deduped

    def _row_to_hit(self, row: sqlite3.Row, query: str) -> KnowledgeHit:
        score = float(row["score"])
        try:
            mtime_str = row["mtime"]
            if mtime_str:
                from datetime import datetime, UTC
                file_dt = datetime.fromisoformat(mtime_str.replace("Z", "+00:00"))
                days_old = (datetime.now(UTC) - file_dt).days
                if days_old > 30:
                    score += min(3.0, (days_old - 30) / 30.0)
        except (KeyError, ValueError, TypeError):
            pass
        haystack = f"{row['path']} {row['title']}".lower()
        query_l = query.lower()
        for product, aliases in PRODUCT_ALIASES.items():
            if any(alias in query_l for alias in aliases) and any(alias in haystack for alias in aliases):
                score -= 10
        if row["area"] == "fact" and re.search(r"\b(lỗi|bug|issue|ticket|sự cố|error|fail)\b", query_l):
            score -= 5
        score += self._path_score_adjustment(query, row["path"], row["title"], row["area"]) * 20
        if row["area"] == "knowledge" and not self._has_non_knowledge_intent(query):
            score -= 5
        snippet = compact_text(str(row["content"]))[:400]
        return KnowledgeHit(
            path=row["path"],
            title=row["title"],
            snippet=snippet,
            score=score,
            lines=f"L{row['start_line']}-L{row['end_line']}",
            product=row["product"],
            area=row["area"],
        )

    @staticmethod
    def _fold(text: str) -> str:
        normalized = unicodedata.normalize("NFD", text.lower())
        return "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")

    @classmethod
    def _has_any(cls, folded_text: str, markers: list[str]) -> bool:
        return any(cls._contains_phrase(folded_text, marker) for marker in markers)

    @staticmethod
    def _contains_phrase(folded_text: str, marker: str) -> bool:
        folded_marker = KBService._fold(marker)
        if re.search(r"[a-z0-9]", folded_marker):
            return re.search(rf"(?<![a-z0-9]){re.escape(folded_marker)}(?![a-z0-9])", folded_text) is not None
        return folded_marker in folded_text

    @classmethod
    def _query_products(cls, query: str) -> list[str]:
        q = cls._fold(query)
        products: list[str] = []
        for product, aliases in PRODUCT_ALIASES.items():
            markers = [product, *aliases]
            if cls._has_any(q, markers):
                products.append(product)
        return products

    @classmethod
    def _has_non_knowledge_intent(cls, query: str) -> bool:
        q = cls._fold(query)
        markers = [
            "strategy",
            "chien luoc",
            "roadmap",
            "ke hoach",
            "plan",
            "ap26",
            "kpi",
            "okr",
            "metric",
            "chi so",
            "monthly report",
            "highlight",
            "ship",
            "ticket",
            "issue",
            "transid",
            "ma loi",
            "error",
            "fail",
            "source code",
            "ma nguon",
            "ham",
            "function",
            "file nao",
            "checklist",
            "audit",
            "financial safety",
            "trust-building",
        ]
        return cls._has_any(q, markers)

    @classmethod
    def _path_score_adjustment(cls, query: str, path: str, title: str, area: str) -> float:
        q = cls._fold(query)
        p = cls._fold(f"{path} {title}")
        delta = 0.0
        query_products = cls._query_products(query)

        for product in query_products:
            markers = [product, *PRODUCT_ALIASES.get(product, [])]
            if cls._has_any(p, markers):
                delta -= 0.25
            elif any(cls._has_any(p, [other, *aliases]) for other, aliases in PRODUCT_ALIASES.items() if other != product):
                delta += 0.2

        exact_issues = [cls._fold(issue) for issue in re.findall(r"\bISSUE-\d+\b", query, flags=re.IGNORECASE)]
        if exact_issues and path.startswith("03. Fact/CS Ticket/"):
            if any(issue in p for issue in exact_issues):
                delta -= 1.2
            else:
                delta += 0.9

        strategy_like = cls._has_any(q, ["strategy", "chien luoc", "roadmap", "uu tien", "priority", "ap26", "focus"])
        if not strategy_like and cls._contains_phrase(q, "wealth") and cls._has_any(q, ["san pham", "products", "liet ke", "hien co"]):
            if path == "05. Knowledge/_Index.md":
                delta -= 0.9

        if strategy_like:
            if path.startswith("01. Objective/Product Strategy/"):
                delta -= 1.0
            if path.startswith("02. Context/Confluence/Wealth General/"):
                delta -= 0.35
            if "plan" in p or "roadmap" in p or "planning" in p:
                delta -= 0.15

        if cls._has_any(q, ["monthly report", "highlight", "ship", "thang 2", "thang 3", "thang 4", "thang 5", "feb", "mar", "apr", "may"]):
            if path.startswith("01. Objective/Monthly Report/"):
                delta -= 0.9
            month_map = {
                "thang 2": ["feb"],
                "thang 3": ["mar"],
                "thang 4": ["apr"],
                "thang 5": ["may"],
            }
            for marker, aliases in month_map.items():
                if cls._contains_phrase(q, marker) and any(alias in p for alias in aliases):
                    delta -= 0.25

        if cls._has_any(q, ["kpi", "okr", "objective", "kr", "product contribution", "shared kpi", "chi so"]):
            if path.startswith("01. Objective/Product KPI/"):
                delta -= 0.85
            if cls._contains_phrase(q, "product contribution") and "product contribution" in p:
                delta -= 0.25
            if cls._contains_phrase(q, "shared kpi") and "shared kpi" in p:
                delta -= 0.25
            if cls._contains_phrase(q, "okr") and "okr" in p:
                delta -= 0.25

        if cls._has_any(q, ["checklist", "audit", "financial safety", "trust-building", "trust building"]):
            if path.startswith(".agents/skills/Audit/Product Audit/") or path.startswith("04. Skill/"):
                delta -= 0.9
            if path.startswith("01. Objective/Product KPI/03. Product Audit/"):
                delta -= 0.45

        if cls._has_any(q, ["issue", "ticket", "loi", "error", "fail", "khieu nai", "cs"]):
            if path.startswith("03. Fact/CS Ticket/"):
                delta -= 0.75
            if path.startswith("03. Fact/Issue Investigation/"):
                delta -= 0.65
            if path.startswith("02. Context/Jira/"):
                delta -= 0.35

        if cls._has_any(q, ["transid", "transaction", "timeout", "retry"]):
            if path.startswith("03. Fact/Issue Investigation/"):
                delta -= 0.9
            if path.startswith("03. Fact/Source Code/"):
                delta -= 0.35

        source_like = cls._has_any(q, ["source code", "ma nguon", "code", "ham", "function", "logic", "file nao", "controller", "handler"])
        if source_like:
            if path.startswith("03. Fact/Source Code/"):
                delta -= 0.9
                if cls._contains_phrase(q, "redemption") and "redemption" in p:
                    delta -= 0.55
                if any(marker in p for marker in ["controller", "handler", "service"]):
                    delta -= 0.2
                if any(marker in p for marker in ["/statics/", ".json", "nested-divisions"]):
                    delta += 0.45

        if cls._has_any(q, ["fs profile", "kyc", "risk assessment"]):
            if path.startswith("02. Context/Confluence/FS Profile/") or path.startswith("02. Context/Confluence/FS Profle/"):
                delta -= 0.8
            if path.startswith("02. Context/Confluence/FS Hub/"):
                delta -= 0.35

        if area == "knowledge" and cls._has_non_knowledge_intent(query):
            delta += 0.12
        return delta

    def _safe_current_path(self, rel_path: str) -> Path:
        if not self.current_link.exists():
            raise FileNotFoundError("KB current is not available")
        root = self.current_link.resolve()
        candidate = (root / rel_path).resolve()
        try:
            candidate.relative_to(root)
        except ValueError as exc:
            raise ValueError("path traversal blocked") from exc
        return candidate

    def _grep_rg(self, rg: str, pattern: str, root: Path, max_results: int) -> str:
        cmd = [rg, "--line-number", "--no-heading", "--color", "never", pattern, str(root)]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=10, check=False)
        except subprocess.TimeoutExpired:
            return '{"error":"grep timeout"}'
        if proc.returncode not in {0, 1}:
            return safe_json({"error": "grep failed"})
        rows: list[dict[str, Any]] = []
        kb_root = self.current_link.resolve()
        for line in proc.stdout.splitlines():
            if len(rows) >= max_results:
                break
            parts = line.split(":", 2)
            if len(parts) < 3:
                continue
            path_s, line_no, text = parts
            try:
                rel = Path(path_s).resolve().relative_to(kb_root).as_posix()
            except ValueError:
                continue
            rows.append({"path": rel, "line_no": int(line_no), "line": text[:500]})
        return safe_json(rows)

    def _grep_python(self, pattern: str, root: Path, max_results: int) -> str:
        try:
            regex = re.compile(pattern)
        except re.error:
            regex = re.compile(re.escape(pattern))
        rows: list[dict[str, Any]] = []
        kb_root = self.current_link.resolve()
        for path in root.rglob("*"):
            if len(rows) >= max_results:
                break
            if not path.is_file() or path.suffix.lower() not in TEXT_EXTENSIONS:
                continue
            for idx, line in enumerate(self._read_text(path).splitlines(), start=1):
                if regex.search(line):
                    rows.append({"path": path.relative_to(kb_root).as_posix(), "line_no": idx, "line": line[:500]})
                    if len(rows) >= max_results:
                        break
        return safe_json(rows)

    def _cleanup_old_versions(self) -> None:
        keep = max(1, self.settings.kb_keep_versions)
        with self.db.connect() as conn:
            rows = conn.execute("SELECT id FROM kb_versions ORDER BY id DESC").fetchall()
        for row in rows[keep:]:
            version_id = int(row["id"])
            path = self.versions_dir / str(version_id)
            if path.exists():
                shutil.rmtree(path)
            with self.db.connect() as conn:
                conn.execute("DELETE FROM kb_files WHERE kb_version=?", (version_id,))
                conn.execute("DELETE FROM kb_versions WHERE id=?", (version_id,))
                conn.commit()
        self._gc_chunks()

    def _gc_chunks(self) -> None:
        with self.db.connect() as conn:
            conn.execute(
                """
                DELETE FROM chunks_fts
                WHERE rowid IN (
                    SELECT rowid_fts FROM chunks_meta
                    WHERE NOT EXISTS (
                        SELECT 1 FROM kb_files
                        WHERE kb_files.path = chunks_meta.path
                          AND kb_files.sha256 = chunks_meta.file_sha
                    )
                )
                """
            )
            conn.execute(
                """
                DELETE FROM chunks_meta
                WHERE NOT EXISTS (
                    SELECT 1 FROM kb_files
                    WHERE kb_files.path = chunks_meta.path
                      AND kb_files.sha256 = chunks_meta.file_sha
                )
                """
            )
            conn.commit()

    def _hardlink_tree(self, src: Path, dst: Path) -> None:
        if dst.exists():
            shutil.rmtree(dst)
        dst.mkdir(parents=True, exist_ok=True)
        for item in src.rglob("*"):
            rel = item.relative_to(src)
            target = dst / rel
            if item.is_dir():
                target.mkdir(parents=True, exist_ok=True)
            elif item.is_file():
                target.parent.mkdir(parents=True, exist_ok=True)
                try:
                    os.link(item, target)
                except OSError:
                    shutil.copy2(item, target)

    async def apply_delta(self, archive_path: Path, meta: dict[str, Any], uploaded_by: str = "sync-api") -> int:
        self.ensure_dirs()
        active_id = self.active_version()
        if not active_id:
            raise ValueError("No active version to apply delta on")

        base_version = int(meta["base_version"])
        client_host = meta.get("client_host", "unknown")
        deleted = meta.get("deleted", [])
        with self.db.connect() as conn:
            active_count = conn.execute("SELECT COUNT(*) FROM kb_files WHERE kb_version=?", (active_id,)).fetchone()[0]
        if active_count > 10 and len(deleted) > 0.3 * active_count:
            raise ValueError(f"Deletions exceed 30% safety limit: deleting {len(deleted)} of {active_count} files")
        added_modified = meta.get("added_modified", [])
        client_manifest_sha = meta.get("client_manifest_sha")

        version_id = self._create_version(uploaded_by, archive_path.name)
        with self.db.connect() as conn:
            conn.execute(
                "UPDATE kb_versions SET kind='delta', base_version=? WHERE id=?",
                (base_version, version_id),
            )
            conn.commit()

        target = self.versions_dir / str(version_id)

        try:
            active_path = self.versions_dir / str(active_id)
            self._hardlink_tree(active_path, target)

            for rel_path in deleted:
                file_to_del = target / rel_path
                if file_to_del.exists() and file_to_del.is_file():
                    file_to_del.unlink()
                parent = file_to_del.parent
                while parent != target:
                    if parent.exists() and not any(parent.iterdir()):
                        parent.rmdir()
                        parent = parent.parent
                    else:
                        break

            max_uncompressed = self.settings.kb_delta_max_mb * 1024 * 1024 * 5
            total = 0
            with zipfile.ZipFile(archive_path) as zf:
                infos = zf.infolist()
                for info in infos:
                    rel = Path(info.filename)
                    if info.is_dir():
                        continue
                    if rel.is_absolute() or ".." in rel.parts:
                        raise ValueError("Zip path traversal blocked")
                    if self._excluded_zip_path(rel):
                        continue
                    total += info.file_size
                    if total > max_uncompressed:
                        raise ValueError("Zip uncompressed size exceeds safety limit")

                    destination = (target / rel).resolve()
                    try:
                        destination.relative_to(target.resolve())
                    except ValueError as exc:
                        raise ValueError("Zip path traversal blocked") from exc

                    destination.parent.mkdir(parents=True, exist_ok=True)
                    if destination.exists() or destination.is_symlink():
                        destination.unlink()

                    tmp_dest = destination.parent / f"{destination.name}.tmp-{uuid.uuid4().hex}"
                    with zf.open(info, "r") as src, tmp_dest.open("wb") as dst:
                        shutil.copyfileobj(src, dst, length=1024 * 1024)
                    os.replace(tmp_dest, destination)

            old_manifest = {}
            with self.db.connect() as conn:
                rows = conn.execute(
                    "SELECT path, sha256, size, mtime FROM kb_files WHERE kb_version=?",
                    (active_id,),
                ).fetchall()
                for row in rows:
                    old_manifest[row["path"]] = {
                        "sha256": row["sha256"],
                        "size": row["size"],
                        "mtime": row["mtime"],
                    }

            new_manifest = {}
            for path, meta_val in old_manifest.items():
                if path in deleted:
                    continue
                if path in added_modified:
                    continue
                new_manifest[path] = meta_val

            now_str = utc_now()
            stats = {"files": 0, "chunks": 0, "bytes": 0}

            with self.db.connect() as conn:
                for rel_path in added_modified:
                    file_path = target / rel_path
                    if not file_path.exists() or not file_path.is_file():
                        continue
                    size = file_path.stat().st_size
                    sha = sha256_file(file_path)
                    new_manifest[rel_path] = {
                        "sha256": sha,
                        "size": size,
                        "mtime": now_str,
                    }

                    if file_path.suffix.lower() not in TEXT_EXTENSIONS:
                        continue
                    if size > self.settings.kb_file_max_mb * 1024 * 1024:
                        continue

                    exists = conn.execute(
                        "SELECT 1 FROM chunks_meta WHERE path=? AND file_sha=? LIMIT 1",
                        (rel_path, sha),
                    ).fetchone()
                    if exists:
                        continue

                    text = self._read_text(file_path)
                    title = self._title_for(rel_path, text)
                    product = self._classify_product(rel_path)
                    area = self._classify_area(rel_path)
                    for chunk, start_line, end_line in self._chunk_text(text, rel_path):
                        cur = conn.execute(
                            "INSERT INTO chunks_fts(path, title, product, area, content) VALUES (?,?,?,?,?)",
                            (rel_path, title, product, area, chunk),
                        )
                        conn.execute(
                            "INSERT INTO chunks_meta(rowid_fts, path, file_sha, start_line, end_line, mtime) VALUES (?,?,?,?,?,?)",
                            (cur.lastrowid, rel_path, sha, start_line, end_line, now_str),
                        )
                        stats["chunks"] += 1

                for path, meta_val in new_manifest.items():
                    conn.execute(
                        """
                        INSERT OR REPLACE INTO kb_files(kb_version, path, sha256, size, mtime)
                        VALUES (?,?,?,?,?)
                        """,
                        (version_id, path, meta_val["sha256"], meta_val["size"], meta_val["mtime"]),
                    )

                skills = workflows = 0
                registry_modified = any(
                    p.startswith(".agents/skills/") or p.startswith(".agents/workflows/")
                    for p in deleted + added_modified
                )

                if client_manifest_sha:
                    sorted_files = {k: {"sha256": v["sha256"], "size": v["size"]} for k, v in sorted(new_manifest.items())}
                    manifest_json = json.dumps(sorted_files, separators=(",", ":"), sort_keys=True)
                    manifest_sha = hashlib.sha256(manifest_json.encode("utf-8")).hexdigest()
                    if manifest_sha != client_manifest_sha:
                        raise ValueError(
                            f"Manifest checksum mismatch. Expected: {client_manifest_sha}, got: {manifest_sha}"
                        )

                conn.commit()

            total_files = len(new_manifest)
            total_bytes = sum(x["size"] for x in new_manifest.values())

            activate_policy = self.settings.kb_sync_activate
            should_activate = activate_policy == "auto"

            if should_activate:
                self.activate_version(version_id)
                if registry_modified:
                    skills, workflows = self.reload_registries()
            else:
                with self.db.connect() as conn:
                    conn.execute(
                        """
                        UPDATE kb_versions
                        SET status='ready', file_count=?, chunk_count=?, total_bytes=?, skill_count=?, workflow_count=?
                        WHERE id=?
                        """,
                        (total_files, stats["chunks"], total_bytes, skills, workflows, version_id),
                    )
                    conn.commit()

            change_summary = {
                "added": len([p for p in added_modified if p not in old_manifest]),
                "modified": len([p for p in added_modified if p in old_manifest]),
                "deleted": len(deleted),
            }

            with self.db.connect() as conn:
                conn.execute(
                    """
                    INSERT OR REPLACE INTO sync_state(id, last_sync_at, last_client_host, last_result, last_error)
                    VALUES (1, ?, ?, ?, NULL)
                    """,
                    (
                        now_str,
                        client_host,
                        f"v{version_id} (active)" if should_activate else f"v{version_id} (ready)",
                    ),
                )
                conn.execute(
                    """
                    UPDATE kb_versions
                    SET change_summary=?
                    WHERE id=?
                    """,
                    (json.dumps(change_summary), version_id),
                )
                conn.commit()

            if self.telegram_client and self.telegram_client.configured and self.settings.telegram_owner_user_ids:
                added_count = change_summary["added"]
                modified_count = change_summary["modified"]
                deleted_count = change_summary["deleted"]

                if should_activate:
                    msg = (
                        f"🔄 <b>KB sync thành công</b>\n"
                        f"Thêm: {added_count}, Sửa: {modified_count}, Xóa: {deleted_count} file\n"
                        f"Version mới: <b>v{version_id} (Active)</b>\n"
                        f"Nguồn: {client_host}"
                    )
                else:
                    msg = (
                        f"🔄 <b>KB sync chờ duyệt</b>\n"
                        f"Thêm: {added_count}, Sửa: {modified_count}, Xóa: {deleted_count} file\n"
                        f"Version mới: <b>v{version_id} (Chờ duyệt)</b>\n"
                        f"Duyệt bằng lệnh: <code>/kb_activate {version_id}</code>\n"
                        f"Nguồn: {client_host}"
                    )
                for owner_id in self.settings.telegram_owner_user_ids:
                    try:
                        await self.telegram_client.send_message(owner_id, msg)
                    except Exception as e:
                        log_event("error", "telegram_notify_failed", error=str(e), owner_id=owner_id)

            if self.audit:
                self.audit.record("sync-api", "kb_sync_success", str(version_id), change_summary)

            return version_id

        except Exception as exc:
            with self.db.connect() as conn:
                conn.execute(
                    """
                    INSERT OR REPLACE INTO sync_state(id, last_sync_at, last_client_host, last_result, last_error)
                    VALUES (1, ?, ?, 'failed', ?)
                    """,
                    (utc_now(), client_host, str(exc)),
                )
                conn.execute("UPDATE kb_versions SET status='failed', error=? WHERE id=?", (str(exc), version_id))
                conn.commit()
            if self.audit:
                self.audit.record("sync-api", "kb_sync_failed", str(version_id), {"error": str(exc)})
            raise

    @staticmethod
    def _read_text(path: Path) -> str:
        ext = path.suffix.lower()
        if ext == ".pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(path)
                text_parts = []
                for idx, page in enumerate(reader.pages, start=1):
                    page_text = page.extract_text() or ""
                    text_parts.append(f"--- Page {idx} ---\n{page_text}")
                return "\n\n".join(text_parts)
            except Exception as e:
                return f"[Lỗi đọc PDF {path.name}: {e}]"
        elif ext == ".docx":
            try:
                import docx
                doc = docx.Document(path)
                text_parts = []
                for para in doc.paragraphs:
                    if para.text.strip():
                        text_parts.append(para.text)
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text.strip() for cell in row.cells]
                        text_parts.append(" | ".join(row_text))
                return "\n".join(text_parts)
            except Exception as e:
                return f"[Lỗi đọc DOCX {path.name}: {e}]"
        elif ext in (".xlsx", ".xlsm"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                text_parts = []
                for sheet_name in wb.sheetnames:
                    text_parts.append(f"--- Sheet: {sheet_name} ---")
                    sheet = wb[sheet_name]
                    for row in sheet.iter_rows(values_only=True):
                        if any(row):
                            row_str = " | ".join(str(val) if val is not None else "" for val in row)
                            text_parts.append(row_str)
                return "\n".join(text_parts)
            except Exception as e:
                return f"[Lỗi đọc XLSX {path.name}: {e}]"

        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        return path.read_text(encoding="utf-8", errors="ignore")

    @staticmethod
    def _match_query(query: str) -> str:
        tokens = re.findall(r"[\wÀ-ỹ]+", query.lower(), flags=re.UNICODE)
        filtered = []
        for token in tokens:
            if len(token) < 2:
                continue
            if token.isdigit() and len(token) > 6:
                continue
            filtered.append(token)
            if len(filtered) >= 12:
                break

        # Vietnamese-to-English technical mappings
        extra_tokens = []
        lower_query = query.lower()
        mappings = {
            "nạp tiền": ["deposit"],
            "nap tien": ["deposit"],
            "rút tiền": ["withdraw", "redemption", "redeem"],
            "rut tien": ["withdraw", "redemption", "redeem"],
            "thanh toán": ["payment"],
            "thanh toan": ["payment"],
            "đối soát": ["reconciliation", "reconcile"],
            "doi soat": ["reconciliation", "reconcile"],
            "chuyển tiền": ["transfer"],
            "chuyen tien": ["transfer"],
            "tài khoản": ["account"],
            "tai khoan": ["account"],
            "số dư": ["balance"],
            "so du": ["balance"],
            "kỹ thuật": ["controller", "handler", "service", "internal"],
            "ky thuat": ["controller", "handler", "service", "internal"],
            "mã nguồn": ["controller", "handler", "service", "internal"],
            "ma nguon": ["controller", "handler", "service", "internal"],
            "code": ["controller", "handler", "service", "internal"],
            "lỗi": ["error", "fail", "failed", "pending"],
            "loi": ["error", "fail", "failed", "pending"],
            "thất bại": ["fail", "failed"],
            "that bai": ["fail", "failed"],
        }
        for vn_term, eng_terms in mappings.items():
            if vn_term in lower_query:
                extra_tokens.extend(eng_terms)

        all_tokens = filtered + list(set(extra_tokens))
        return " OR ".join(all_tokens)

    @staticmethod
    def _title_for(rel: str, text: str) -> str:
        for line in text.splitlines()[:40]:
            if line.startswith("#"):
                title = line.lstrip("#").strip()
                if title:
                    return title[:200]
        return rel

    @staticmethod
    def _chunk_text(text: str, path: str = "", max_chars: int = 1400, overlap: int = 180) -> list[tuple[str, int, int]]:
        lines = text.splitlines()
        line_headings = []
        h1, h2, h3 = None, None, None
        for idx, line in enumerate(lines, start=1):
            stripped = line.strip()
            if stripped.startswith("# "):
                h1 = stripped[2:].strip()
                h2, h3 = None, None
            elif stripped.startswith("## "):
                h2 = stripped[3:].strip()
                h3 = None
            elif stripped.startswith("### "):
                h3 = stripped[4:].strip()
            line_headings.append((h1, h2, h3))

        paragraphs: list[tuple[str, int, int, str]] = []
        buf: list[str] = []
        start = 1
        for idx, line in enumerate(lines, start=1):
            if line.strip():
                if not buf:
                    start = idx
                buf.append(line)
            elif buf:
                h1_val, h2_val, h3_val = line_headings[start - 1]
                parts = []
                if path:
                    parts.append(os.path.basename(path))
                if h1_val:
                    parts.append(h1_val)
                if h2_val:
                    parts.append(h2_val)
                if h3_val:
                    parts.append(h3_val)
                bc = " > ".join(parts)
                paragraphs.append(("\n".join(buf), start, idx - 1, bc))
                buf = []
        if buf:
            h1_val, h2_val, h3_val = line_headings[start - 1]
            parts = []
            if path:
                parts.append(os.path.basename(path))
            if h1_val:
                parts.append(h1_val)
            if h2_val:
                parts.append(h2_val)
            if h3_val:
                parts.append(h3_val)
            bc = " > ".join(parts)
            paragraphs.append(("\n".join(buf), start, start + len(buf) - 1, bc))

        chunks: list[tuple[str, int, int]] = []
        current = ""
        current_start = 1
        current_end = 1
        current_bc = ""

        for para, start_line, end_line, bc in paragraphs:
            prefix = f"[Breadcrumb: {bc}]\n\n" if bc else ""
            if len(prefix + para) > max_chars:
                if current:
                    chunks.append((f"[Breadcrumb: {current_bc}]\n\n" + current if current_bc else current, current_start, current_end))
                    current = ""
                pos = 0
                while pos < len(para):
                    part = para[pos : pos + max_chars - len(prefix)]
                    chunks.append((prefix + part.strip() if bc else part.strip(), start_line, end_line))
                    pos += max_chars - len(prefix) - overlap
                continue

            if not current:
                current = para
                current_start = start_line
                current_end = end_line
                current_bc = bc
            elif len(current) + len(para) + 2 <= max_chars - len(prefix):
                current += "\n\n" + para
                current_end = end_line
            else:
                chunks.append((f"[Breadcrumb: {current_bc}]\n\n" + current if current_bc else current, current_start, current_end))
                current = para
                current_start = start_line
                current_end = end_line
                current_bc = bc

        if current:
            chunks.append((f"[Breadcrumb: {current_bc}]\n\n" + current if current_bc else current, current_start, current_end))

        return [chunk for chunk in chunks if chunk[0]]

    @staticmethod
    def _classify_area(rel: str) -> str:
        first = rel.split("/", 1)[0].lower()
        if first.startswith("01."):
            return "objective"
        if first.startswith("02."):
            return "context"
        if first.startswith("03."):
            return "fact"
        if first.startswith("04."):
            return "skill"
        if first.startswith("05."):
            return "knowledge"
        if first == ".agents":
            return "agents"
        return ""

    @staticmethod
    def _classify_product(rel: str) -> str:
        parts = [p.strip().lower() for p in rel.split("/")]
        cleaned_parts = []
        for p in parts:
            p_clean = re.sub(r"^\d+[\.\s-]*", "", p).strip()
            cleaned_parts.append(p_clean)

        for product, aliases in PRODUCT_ALIASES.items():
            aliases_lower = [a.lower() for a in aliases] + [product.lower()]
            for part in cleaned_parts:
                if part in aliases_lower:
                    return product

        lower = rel.lower()
        for product, aliases in PRODUCT_ALIASES.items():
            aliases_lower = [a.lower() for a in aliases] + [product.lower()]
            for alias in aliases_lower:
                pattern = r"(?<![a-zA-Z0-9_À-ỹ])" + re.escape(alias) + r"(?![a-zA-Z0-9_À-ỹ])"
                if re.search(pattern, lower):
                    return product
        return ""
